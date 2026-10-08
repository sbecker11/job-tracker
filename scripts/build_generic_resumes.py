"""Build the generic (not JD-tailored) résumés from the master résumé.

  .venv/bin/python scripts/build_generic_resumes.py            # all variants
  .venv/bin/python scripts/build_generic_resumes.py full_stack

Reads ~/Desktop/Resumes/Details/SHAWN_BECKER_full_resume_indexed.docx, strips
[N]/[skill:] tags, hides education and certification years, and writes
~/Desktop/Resumes/Details/SHAWN_BECKER_<variant>.docx.
"""

from __future__ import annotations

import copy
import re
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Pt

from job_tracker.pipeline.skill_tags import strip_skill_tags_file

DETAILS = Path.home() / "Desktop/Resumes/Details"
MASTER = DETAILS / "SHAWN_BECKER_full_resume_indexed.docx"

VARIANTS = {
    "full_stack": {
        "title": "Full-Stack Web Applications — Vue 3, React, Node.js & Java/Spring Boot",
        "summary": (
            "Senior full-stack engineer (individual contributor) with 30+ years building production software, "
            "from front-end UIs through backend services to the data platforms behind them. Recent client work "
            "includes Vue 3 + Node.js registration, login, and profile flows for NuSkin and Python REST admin "
            "tools for Adobe’s SEO platform; portfolio work includes resume-flyer (Vue 3, TypeScript, Vite, "
            "Express.js) and spexture-com (React 18, Express.js, JWT/RBAC auth, PostgreSQL, Docker). Deep Java "
            "and Spring background from co-founding HomePortfolio (Java, JSP, Oracle) and ClipFile (Java, "
            "Spring, DynamoDB), plus ~2.75 years owning SeniorLink’s AWS data platform. MIT PhD in computer "
            "vision and computer graphics; co-inventor on three NLP patents."
        ),
        "skills": [
            ("Front end", "Vue 3, React 18, TypeScript, Vite, Vuetify, Vuelidate, JSP"),
            ("Back end & APIs", "Node.js, Express.js, Java 21, Spring Boot, Python, FastAPI, REST, JWT/bcrypt auth, RBAC"),
            ("Data", "PostgreSQL, Aurora PostgreSQL, pgvector, Oracle, DynamoDB, Redshift, Snowflake, BigQuery"),
            ("Cloud & delivery", "AWS (ECS Fargate, S3, EMR, Glue, Bedrock, CloudWatch, CodePipeline/CodeBuild/CodeDeploy), "
                                 "GCP, Docker, Docker Compose, CloudFormation, AWS CDK, Artifactory"),
            ("Practices", "Automated testing (linkage-engine: 149 tests, 94.5% coverage), Agile/Scrum (CSM), Jira"),
        ],
        "projects": ["resume-flyer", "spexture-com", "color-palette-maker", "linkage-engine",
                     "ecommerce-semantic-search"],
    },
    "software_engineer": {
        "title": "Java/Spring Boot, Python & AWS Distributed Systems",
        "summary": (
            "Senior software engineer (individual contributor) with 30+ years building production systems in "
            "Java, Python, and AWS, much of it in regulated domains: healthcare, mortgage finance, and insurance "
            "claims. Spent ~2.75 years at SeniorLink (now Careforth) owning an event-driven data platform "
            "(Kinesis and RabbitMQ ingestion, PySpark on EMR, Parquet in S3, Redshift, Tableau) and built dbt "
            "and AWS Glue pipelines with explicit lineage for Fannie Mae. Long Java/Spring track record from "
            "co-founding HomePortfolio and ClipFile, carried forward into linkage-engine (Java 21, Spring Boot, "
            "pgvector on Aurora PostgreSQL, ECS Fargate, 149 tests). MIT PhD in computer vision and computer "
            "graphics; co-inventor on three NLP patents."
        ),
        "skills": [
            ("Languages", "Java 21, Python, PySpark, SQL, TypeScript, C++"),
            ("Frameworks", "Spring Boot, Spring AI, Virtual Threads, FastAPI, Node.js/Express.js, Vue 3, React"),
            ("Data & streaming", "Kafka, Kinesis, RabbitMQ, Airflow, dbt, AWS Glue, EMR, Parquet/S3, Redshift, Snowflake, "
                                 "BigQuery, PostgreSQL/pgvector, Oracle, DynamoDB, Flyway"),
            ("Cloud & delivery", "AWS (ECS Fargate, Aurora PostgreSQL Serverless v2, Bedrock, CloudWatch, "
                                 "CodeCommit/CodeBuild/CodePipeline/CodeDeploy), GCP, CloudFormation, AWS CDK, Docker"),
            ("Security & practices", "CyberArk PAM, mTLS, JWT/RBAC, Xray test management, Agile/Scrum (CSM), Jira"),
        ],
        "projects": ["linkage-engine", "ecommerce-semantic-search", "resume-parser", "resume-flyer",
                     "recruiting-automation"],
    },
    "data_engineer": {
        "title": "Distributed, Event-Driven Data Platforms for Regulated Systems",
        "summary": (
            "Senior data engineer (individual contributor) with 30+ years building production systems, focused "
            "since 2017 on event-driven data platforms in regulated domains: healthcare, mortgage finance, and "
            "insurance. At SeniorLink (now Careforth) spent ~2.75 years owning a JSON schema-governed event "
            "pipeline: Kinesis and RabbitMQ ingestion, PySpark on AWS EMR, date-partitioned Parquet in S3, "
            "Redshift OLAP, Tableau self-service BI, and Flyway migrations. As a consultant, built dbt models and "
            "AWS Glue pipelines with metadata lineage and schema-evolution tracking for Fannie Mae, Kafka "
            "streaming with Airflow orchestration for Cigna, a Segment CDP feeding BigQuery for Warner Bros. "
            "Games, and a Type-2 SCD PostgreSQL warehouse for Greenseed. MIT PhD in computer vision and computer "
            "graphics; co-inventor on three NLP patents."
        ),
        "skills": [
            ("Languages", "Python, PySpark, SQL, Java 21, TypeScript"),
            ("Pipelines & orchestration", "Airflow, dbt, AWS Glue + Glue Crawlers, EMR, Flyway"),
            ("Streaming & messaging", "Kafka, Kinesis, RabbitMQ, Segment CDP"),
            ("Warehousing & BI", "Redshift, Snowflake, BigQuery, S3/Parquet, PostgreSQL (Type-2 SCD), Aurora "
                                 "PostgreSQL, pgvector, Oracle, DynamoDB, Tableau, Looker"),
            ("Cloud & delivery", "AWS (EMR, Glue, S3, ECS Fargate, CloudWatch, "
                                 "CodeCommit/CodeBuild/CodePipeline/CodeDeploy), GCP, CloudFormation, AWS CDK, Docker"),
            ("Security & practices", "CyberArk PAM, mTLS, Xray test management, Agile/Scrum (CSM), Jira"),
        ],
        "projects": ["linkage-engine", "ecommerce-semantic-search", "healthcare-agentic-snowflake-rag",
                     "recruiting-automation", "resume-parser"],
    },
}

SIERRA_VISTA_CLIENTS = (
    "Clients: Rowe International (Lead Engineer/System Architect — Java, JSP, PostgreSQL); Intrusic (System "
    "Architect — C++ packet analysis for network intrusion detection); AMI (Data Architect/Project Manager — "
    "Java, JSP, FileMaker); Eleven LLC (CTO — MS Mobile); Coca-Cola Corp (Technical Lead/Project Manager — "
    "SQL Anywhere, NServiceBus); MSC/One Call ICD/CPT Workers Compensation claims (Senior Software Engineer, "
    "Jacksonville, FL)."
)
SIERRA_VISTA_BULLETS = ("• Rowe International", "• Intrusic", "• AMI", "• Eleven LLC", "• Coca-Cola Corp",
                        "• MSC / One Call")
UNBOLD_PREFIXES = ("~2.75 yrs PySpark", "Java, JSP, Spring, and PostgreSQL", "Built a content curation")
YEAR_RANGE = re.compile(r"\s*\((19|20)\d\d\s*[–-]\s*(19|20)\d\d\)")
CERT_YEAR = re.compile(r":\s*(19|20)\d\d\s*$")


def _set_text(p, text: str) -> None:
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""


def _delete(p) -> None:
    p._p.getparent().remove(p._p)


def _find(doc, text: str):
    return next(p for p in doc.paragraphs if p.text.strip() == text)


def build(variant: str) -> Path:
    cfg = VARIANTS[variant]
    with tempfile.TemporaryDirectory() as tmp:
        clean = Path(tmp) / "master_clean.docx"
        strip_skill_tags_file(MASTER, clean, strip_skills=True, strip_indexes=True)
        doc = Document(clean)

    lines = [p for p in doc.paragraphs if p.text.strip()]
    _set_text(lines[2], cfg["title"])
    _set_text(lines[5], cfg["summary"])

    for p in list(doc.paragraphs):
        t = p.text
        if ("University" in t or "Institute" in t) and YEAR_RANGE.search(t):
            _set_text(p, YEAR_RANGE.sub("", t))
        if "historyies" in t:
            for r in p.runs:
                r.text = r.text.replace("historyies", "histories")
        if t.startswith(UNBOLD_PREFIXES):
            for r in p.runs:
                r.bold = False

    in_projects = False
    for p in list(doc.paragraphs):
        t = p.text.strip()
        if t == "Spexture Portfolio Projects":
            in_projects = True
            continue
        if t == "Spexture Client Engagements":
            in_projects = False
        if in_projects and t.startswith("• ") and not any(t.startswith(f"• {k} ") for k in cfg["projects"]):
            _delete(p)

    sv = [p for p in doc.paragraphs if p.text.strip().startswith(SIERRA_VISTA_BULLETS)]
    _set_text(sv[0], SIERRA_VISTA_CLIENTS)
    for p in sv[1:]:
        _delete(p)

    paras = doc.paragraphs
    cert_head = _find(doc, "Licenses & certifications")
    start = next(i for i, p in enumerate(paras) if p._p is cert_head._p)
    certs = []
    for p in paras[start + 1:]:
        if p.text.strip() == "Key skills":
            break
        if p.text.strip():
            certs.append(p)
    _set_text(certs[0], " · ".join(CERT_YEAR.sub("", c.text).strip() for c in certs))
    certs[0].paragraph_format.left_indent = Pt(0)
    for c in certs[1:]:
        _delete(c)

    skills_head = _find(doc, "Key skills")
    paras = doc.paragraphs
    k = next(i for i, p in enumerate(paras) if p._p is skills_head._p)
    old_skills = [p for p in paras[k + 1:] if p.text.strip()]
    template = old_skills[0]._p
    experience = _find(doc, "Experience")
    experience._p.addprevious(skills_head._p)
    for label, items in cfg["skills"]:
        new = copy.deepcopy(template)
        experience._p.addprevious(new)
        p = next(x for x in doc.paragraphs if x._p is new)
        _set_text(p, f"{label}: {items}")
        p.runs[0].text = f"{label}: "
        p.runs[0].bold = True
        tail = copy.deepcopy(p.runs[0]._r)
        p.runs[0]._r.addnext(tail)
        p.runs[1].text = items
        p.runs[1].bold = False
    for p in old_skills:
        _delete(p)
    experience.paragraph_format.space_before = Pt(10)

    body = doc.element.body
    while True:
        last = [e for e in body if e.tag.endswith("}p")][-1]
        if "".join(last.itertext()).strip():
            break
        body.remove(last)

    for p in doc.paragraphs:
        for r in p.runs:
            if r.font.size is not None and r.font.size.pt == 12.0:
                r.font.size = Pt(10.5)

    out = DETAILS / f"SHAWN_BECKER_{variant}.docx"
    doc.save(out)
    return out


def main(argv: list[str]) -> int:
    for v in argv or list(VARIANTS):
        print("saved", build(v))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
