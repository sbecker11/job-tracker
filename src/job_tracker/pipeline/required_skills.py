"""Required-skills extraction: list every skill a lead's JD and recruiter
messages ask for, store it in `lead_skills`, and write `RequiredSkills.docx`
into the lead's package folder.

Extraction is profile-free on purpose: it records what the market asks for,
not how Shawn compares. The met / gap mapping is a separate, global step over
the canonical skill list, so every lead is judged against the same mapping.
"""

from __future__ import annotations

import hashlib
import json
import re
import sqlite3
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from job_tracker.pipeline.models import utc_now_iso
from job_tracker.scoring.scorer import _dealbreaker_sweep, load_framework

MAX_JD_CHARS = 14_000
MAX_MESSAGE_CHARS = 3_000
MAX_MESSAGES = 4

CATEGORIES = (
    "language",
    "framework",
    "cloud",
    "data",
    "database",
    "ml_ai",
    "devops",
    "architecture",
    "practice",
    "domain",
    "certification",
    "education",
    "soft",
)

# Target roles from ~/CLAUDE.md §2: Senior/Staff SWE, Full Stack, Data, ML/AI.
_ROLE_RE = re.compile(
    r"\b(software|full[- ]?stack|back[- ]?end|front[- ]?end|ui|web|data|analytics|etl|databricks|snowflake|"
    r"machine learning|ml|ai|gen(?:erative)?\s*ai|llm|platform|python|java|cloud|applications?|api|integration)\b"
    r".*\b(engineer|developer|architect|programmer)\b"
    r"|\bengineer,?\s+software\b|\bswe\b|\bmlops\b",
    re.I,
)
_ROLE_EXCLUDE_RE = re.compile(
    r"\b(manager|director|head of|vp|vice president|intern|internship|junior|jr\.?|entry[- ]level|student|"
    r"sales|recruiter|support|qa|test|sdet|analyst|scientist|admin(?:istrator)?|devops|sre|site reliability|"
    r"security|network|firmware|embedded|hardware|mobile|ios|android|salesforce|sap|servicenow|mainframe|"
    r"cobol|workday|dynamics)\b",
    re.I,
)

_ALIASES = {
    "js": "JavaScript",
    "javascript": "JavaScript",
    "ts": "TypeScript",
    "typescript": "TypeScript",
    "postgres": "PostgreSQL",
    "postgresql": "PostgreSQL",
    "k8s": "Kubernetes",
    "gcp": "Google Cloud",
    "google cloud platform": "Google Cloud",
    "amazon web services": "AWS",
    "aws": "AWS",
    "reactjs": "React",
    "react.js": "React",
    "node": "Node.js",
    "nodejs": "Node.js",
    "node.js": "Node.js",
    "vue": "Vue.js",
    "vuejs": "Vue.js",
    "vue.js": "Vue.js",
    "ci/cd": "CI/CD",
    "cicd": "CI/CD",
    "llm": "LLMs",
    "llms": "LLMs",
    "large language models": "LLMs",
    "genai": "Generative AI",
    "gen ai": "Generative AI",
    "generative ai": "Generative AI",
    "rag": "RAG",
    "retrieval-augmented generation": "RAG",
    "retrieval augmented generation": "RAG",
    "rest": "REST APIs",
    "rest api": "REST APIs",
    "rest apis": "REST APIs",
    "restful apis": "REST APIs",
    "microservice": "Microservices",
    "microservices": "Microservices",
    "agile": "Agile/Scrum",
    "scrum": "Agile/Scrum",
    "agile/scrum": "Agile/Scrum",
    "sql": "SQL",
    "golang": "Go",
    "c#": "C#",
    ".net": ".NET",
    "terraform": "Terraform",
    "apache spark": "Spark",
    "spark": "Spark",
    "pyspark": "PySpark",
    "apache kafka": "Kafka",
    "kafka": "Kafka",
    "apache airflow": "Airflow",
    "airflow": "Airflow",
    "ml": "Machine Learning",
    "machine learning": "Machine Learning",
}


def canonical_skill(name: str) -> str:
    """Fold common spellings of one skill onto a single name."""
    cleaned = re.sub(r"\s+", " ", (name or "").strip()).strip(" .,;:")
    return _ALIASES.get(cleaned.lower(), cleaned)


@dataclass
class LeadSkill:
    skill: str
    canonical: str
    category: str
    importance: str  # "required" | "preferred"
    years: float | None
    source: str  # "jd" | "message"
    quote: str


@dataclass
class SkillExtraction:
    job_key: str
    role_family: str
    seniority: str
    years_required: float | None
    engagement: dict[str, Any]
    skills: list[LeadSkill] = field(default_factory=list)
    model: str = ""
    input_tokens: int = 0
    output_tokens: int = 0
    cost_usd: float | None = None
    input_hash: str = ""
    raw: dict[str, Any] = field(default_factory=dict)


def is_wheelhouse(title: str, jd_text: str, framework: dict[str, Any] | None = None) -> bool:
    """A target-role title with no load-bearing §3 dealbreaker in the JD."""
    if not _ROLE_RE.search(title or "") or _ROLE_EXCLUDE_RE.search(title or ""):
        return False
    fw = framework if framework is not None else load_framework()
    return not any(hit.load_bearing for hit in _dealbreaker_sweep(jd_text or "", fw))


def inbound_messages(conn: sqlite3.Connection, job_key: str) -> list[str]:
    rows = conn.execute(
        "SELECT body_text FROM job_conversations WHERE job_key = ? AND direction = 'inbound' "
        "AND coalesce(body_text, '') != '' ORDER BY occurred_at DESC LIMIT ?",
        (job_key, MAX_MESSAGES),
    ).fetchall()
    return [r[0][:MAX_MESSAGE_CHARS] for r in rows]


def input_hash(jd_text: str, messages: list[str]) -> str:
    h = hashlib.sha256((jd_text or "").encode("utf-8"))
    for m in messages:
        h.update(b"\x00" + m.encode("utf-8"))
    return h.hexdigest()[:16]


_SYSTEM_PROMPT = f"""You extract the skills a job asks for, from its job description (JD) and the \
recruiter's messages about it. Record what the employer asks for. Do not judge any candidate.

Output ONLY one JSON object, no prose, no markdown fences:
{{
  "role_family": "software | full_stack | backend | frontend | data | ml_ai | platform | other",
  "seniority": "junior | mid | senior | staff | principal | lead | unknown",
  "years_required": <number or null: overall years of experience asked for>,
  "engagement": {{"employment": "full_time | contract | contract_to_hire | unknown",
                  "remote": "remote | hybrid | onsite | unknown",
                  "location": "<city/state/region named, or empty>"}},
  "skills": [
    {{"skill": "<as written, short>",
      "canonical": "<standard short name>",
      "category": "{' | '.join(CATEGORIES)}",
      "importance": "required | preferred",
      "years": <number or null>,
      "source": "jd | message",
      "quote": "<the few words from the text that ask for it, max 15 words>"}}
  ]
}}

Rules:
- One entry per distinct skill. Split lists ("Python, Java, or Go" is three entries).
- canonical: the common product or practice name, consistent across jobs: "AWS Lambda", "AWS Glue", \
"Kubernetes", "PostgreSQL", "React", "Spring Boot", "LLMs", "RAG", "CI/CD", "Agile/Scrum", "Microservices". \
Name specific cloud services individually; use "AWS" only when the text just says AWS.
- importance: "preferred" when the text says preferred, nice to have, a plus, bonus, or similar; else "required".
- source: "message" only when the skill appears in a recruiter message and not in the JD.
- Include domain knowledge (e.g. "Healthcare", "Payments"), degrees and certifications. Include soft skills \
only when stated as a requirement, with category "soft".
- Leave out benefits, perks, and company boilerplate. Never invent a skill the text does not ask for.
"""


def _user_prompt(*, company: str, title: str, jd_text: str, messages: list[str]) -> str:
    parts = [f"Company: {company}", f"Title: {title}", "", "JOB DESCRIPTION:", (jd_text or "")[:MAX_JD_CHARS]]
    for i, m in enumerate(messages, 1):
        parts += ["", f"RECRUITER MESSAGE {i}:", m]
    return "\n".join(parts)


def _to_float(value: Any) -> float | None:
    try:
        return float(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def parse_extraction(data: dict[str, Any], *, job_key: str) -> SkillExtraction:
    seen: set[tuple[str, str]] = set()
    skills: list[LeadSkill] = []
    for s in data.get("skills") or []:
        if not isinstance(s, dict) or not str(s.get("skill") or s.get("canonical") or "").strip():
            continue
        canonical = canonical_skill(str(s.get("canonical") or s.get("skill")))
        importance = "preferred" if str(s.get("importance", "")).lower() == "preferred" else "required"
        key = (canonical.lower(), importance)
        if key in seen:
            continue
        seen.add(key)
        category = str(s.get("category") or "").lower()
        skills.append(
            LeadSkill(
                skill=str(s.get("skill") or canonical).strip(),
                canonical=canonical,
                category=category if category in CATEGORIES else "practice",
                importance=importance,
                years=_to_float(s.get("years")),
                source="message" if str(s.get("source", "")).lower() == "message" else "jd",
                quote=str(s.get("quote") or "").strip()[:200],
            )
        )
    # A skill listed as both required and preferred counts once, as required.
    required = {s.canonical.lower() for s in skills if s.importance == "required"}
    skills = [s for s in skills if s.importance == "required" or s.canonical.lower() not in required]
    engagement = data.get("engagement") if isinstance(data.get("engagement"), dict) else {}
    return SkillExtraction(
        job_key=job_key,
        role_family=str(data.get("role_family") or "other"),
        seniority=str(data.get("seniority") or "unknown"),
        years_required=_to_float(data.get("years_required")),
        engagement=engagement,
        skills=skills,
        raw=data,
    )


def extract_required_skills(
    *,
    job_key: str,
    company: str,
    title: str,
    jd_text: str,
    messages: list[str],
    model: str | None = None,
    client=None,
) -> SkillExtraction:
    from job_tracker.pipeline.llm_apply import DEFAULT_MODEL, _call_and_parse_json

    model = model or DEFAULT_MODEL
    user = _user_prompt(company=company, title=title, jd_text=jd_text, messages=messages)
    data, calls = _call_and_parse_json(
        _SYSTEM_PROMPT, user, model=model, client=client, max_tokens=6000, step="required_skills"
    )
    result = parse_extraction(data, job_key=job_key)
    result.model = model
    result.input_tokens = sum(c.input_tokens for c in calls)
    result.output_tokens = sum(c.output_tokens for c in calls)
    costs = [c.cost_usd for c in calls if c.cost_usd is not None]
    result.cost_usd = sum(costs) if costs else None
    result.input_hash = input_hash(jd_text, messages)
    return result


def predicted_cost(jd_text: str, messages: list[str], *, model: str | None = None) -> float | None:
    from job_tracker.pipeline.llm_apply import DEFAULT_MODEL, _cost_usd, _estimate_tokens

    est_in = _estimate_tokens(_SYSTEM_PROMPT) + _estimate_tokens((jd_text or "")[:MAX_JD_CHARS])
    est_in += sum(_estimate_tokens(m) for m in messages)
    return _cost_usd(model or DEFAULT_MODEL, est_in, 1500)


def stored_hash(conn: sqlite3.Connection, job_key: str) -> str | None:
    row = conn.execute("SELECT input_hash FROM lead_skill_extractions WHERE job_key = ?", (job_key,)).fetchone()
    return row[0] if row else None


def save_extraction(conn: sqlite3.Connection, result: SkillExtraction) -> None:
    conn.execute("DELETE FROM lead_skills WHERE job_key = ?", (result.job_key,))
    conn.executemany(
        "INSERT INTO lead_skills (job_key, skill, canonical, category, importance, years, source, quote) "
        "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
        [
            (result.job_key, s.skill, s.canonical, s.category, s.importance, s.years, s.source, s.quote)
            for s in result.skills
        ],
    )
    conn.execute(
        "INSERT OR REPLACE INTO lead_skill_extractions (job_key, input_hash, model, extracted_at, input_tokens, "
        "output_tokens, cost_usd, role_family, seniority, years_required, engagement_json, raw_json) "
        "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (
            result.job_key,
            result.input_hash,
            result.model,
            utc_now_iso(),
            result.input_tokens,
            result.output_tokens,
            result.cost_usd,
            result.role_family,
            result.seniority,
            result.years_required,
            json.dumps(result.engagement),
            json.dumps(result.raw),
        ),
    )
    conn.commit()


def _fmt_years(value: float | None) -> str:
    if value is None:
        return ""
    return f"{value:g}+"


def render_required_skills_docx(result: SkillExtraction, *, company: str, title: str, out_path: Path) -> Path:
    from docx import Document

    from job_tracker.pipeline.letter_style import apply_template_styles
    from job_tracker.pipeline.llm_apply import _add_table

    doc = Document()
    apply_template_styles(doc)
    doc.add_heading(f"Required skills — {company} — {title}", level=1)
    eng = result.engagement or {}
    facts = [
        f"Role family: {result.role_family}",
        f"Seniority: {result.seniority}",
        f"Years asked: {_fmt_years(result.years_required) or 'not stated'}",
        f"Employment: {eng.get('employment') or 'unknown'}",
        f"Remote: {eng.get('remote') or 'unknown'}" + (f" ({eng['location']})" if eng.get("location") else ""),
    ]
    doc.add_paragraph(" · ".join(facts))
    for importance, heading in (("required", "Required"), ("preferred", "Preferred / nice to have")):
        rows = [
            (s.canonical, s.category, _fmt_years(s.years), "Recruiter message" if s.source == "message" else "JD", s.quote)
            for s in sorted(
                (s for s in result.skills if s.importance == importance), key=lambda s: (s.category, s.canonical.lower())
            )
        ]
        doc.add_heading(f"{heading} ({len(rows)})", level=2)
        if rows:
            _add_table(doc, headers=("Skill", "Category", "Years", "Source", "Quote"), rows=rows)
        else:
            doc.add_paragraph("None listed.")
    doc.add_paragraph(f"Extracted {utc_now_iso()[:10]} by {result.model} from the JD and recruiter messages.")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path
