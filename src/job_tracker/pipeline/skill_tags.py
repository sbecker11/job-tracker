"""Unwrap `[skill:Name]` markers in DOCX text, leaving the inner name."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document

SKILL_TAG_RE = re.compile(r"\[skill:([^\]]+)\]")
# `[1]`, `[1.1]`, `[5.2]`, `[1.1.10]` — not `[skill:…]`
INDEX_TAG_RE = re.compile(r"\[\d+(?:\.\d+)*\]\s*")


def unwrap_skill_tags(text: str) -> str:
    """Turn `[skill:Java]` into `Java`. Leaves `[1]` / `[1.2.3]` indexes alone."""
    return SKILL_TAG_RE.sub(r"\1", text)


def strip_index_tags(text: str) -> str:
    """Remove `[1]` / `[1.2.3]` job-index markers (and a following space)."""
    return INDEX_TAG_RE.sub("", text)


def transform_text(
    text: str, *, strip_skills: bool = True, strip_indexes: bool = False
) -> str:
    if strip_skills:
        text = unwrap_skill_tags(text)
    if strip_indexes:
        text = strip_index_tags(text)
    return text


def _set_paragraph_text(paragraph, text: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def _iter_paragraphs(doc: Document):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def count_tags(text: str) -> tuple[int, int]:
    return len(SKILL_TAG_RE.findall(text)), len(INDEX_TAG_RE.findall(text))


def strip_skill_tags_in_document(
    doc: Document, *, strip_skills: bool = True, strip_indexes: bool = False
) -> tuple[int, int]:
    """Unwrap selected tag classes. Returns (skills, indexes) removed."""
    skills = 0
    indexes = 0
    for para in _iter_paragraphs(doc):
        original = para.text
        new = transform_text(
            original, strip_skills=strip_skills, strip_indexes=strip_indexes
        )
        if new != original:
            s0, i0 = count_tags(original)
            s1, i1 = count_tags(new)
            skills += s0 - s1
            indexes += i0 - i1
            _set_paragraph_text(para, new)
    return skills, indexes


def strip_skill_tags_file(
    src: Path,
    dest: Path | None = None,
    *,
    strip_skills: bool = True,
    strip_indexes: bool = False,
) -> tuple[int, int]:
    """Load `src`, unwrap tags, save to `dest` (default: overwrite `src`)."""
    src = Path(src).expanduser().resolve()
    dest = Path(dest).expanduser().resolve() if dest is not None else src
    doc = Document(src)
    counts = strip_skill_tags_in_document(
        doc, strip_skills=strip_skills, strip_indexes=strip_indexes
    )
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)
    return counts
