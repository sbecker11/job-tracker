from pathlib import Path

from docx import Document

from job_tracker.scripts.index_stripper import _strip_flags, main as cli_main
from job_tracker.pipeline.skill_tags import (
    strip_index_tags,
    strip_skill_tags_file,
    unwrap_skill_tags,
)


def test_unwrap_keeps_index_numbers():
    src = "[1.1.4] • sushi-rag-mcp-server — [skill:FastMCP] + [skill:RAG]"
    assert unwrap_skill_tags(src) == "[1.1.4] • sushi-rag-mcp-server — FastMCP + RAG"


def test_strip_index_tags():
    src = "[1.1.4] • sushi-rag-mcp-server — FastMCP + RAG"
    assert strip_index_tags(src) == "• sushi-rag-mcp-server — FastMCP + RAG"
    assert strip_index_tags("[1] Spexture") == "Spexture"


def test_strip_flags_neither_means_both():
    assert _strip_flags(False, False) == (True, True)


def test_strip_flags_independent():
    assert _strip_flags(True, False) == (True, False)
    assert _strip_flags(False, True) == (False, True)
    assert _strip_flags(True, True) == (True, True)


def test_strip_skill_tags_file(tmp_path: Path):
    src = tmp_path / "in.docx"
    dest = tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("[1] Spexture — [skill:Java] and [skill:Spring Boot]")
    doc.save(src)

    n_skills, n_indexes = strip_skill_tags_file(src, dest)
    assert n_skills == 2
    assert n_indexes == 0
    out = Document(dest)
    assert out.paragraphs[0].text == "[1] Spexture — Java and Spring Boot"
    assert "[skill:" not in out.paragraphs[0].text


def test_strip_indexes_only(tmp_path: Path):
    src = tmp_path / "in.docx"
    dest = tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("[1.2.1] • Adobe — [skill:Python]")
    doc.save(src)

    n_skills, n_indexes = strip_skill_tags_file(
        src, dest, strip_skills=False, strip_indexes=True
    )
    assert n_skills == 0
    assert n_indexes == 1
    assert Document(dest).paragraphs[0].text == "• Adobe — [skill:Python]"


def test_strip_skill_tags_file_with_indexes(tmp_path: Path):
    src = tmp_path / "in.docx"
    dest = tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("[1.2.1] • Adobe — [skill:Python]")
    doc.save(src)

    n_skills, n_indexes = strip_skill_tags_file(src, dest, strip_indexes=True)
    assert n_skills == 1
    assert n_indexes == 1
    assert Document(dest).paragraphs[0].text == "• Adobe — Python"


def test_cli_dry_run_no_flags_counts_both(tmp_path: Path, capsys):
    src = tmp_path / "in.docx"
    doc = Document()
    doc.add_paragraph("[1.1] Spexture — [skill:Java]")
    doc.save(src)
    assert cli_main([str(src), "--dry-run"]) == 0
    out = capsys.readouterr().out
    assert "1 [skill:…] tag(s), 1 index tag(s)" in out


def test_cli_dry_run_skill_only(tmp_path: Path, capsys):
    src = tmp_path / "in.docx"
    doc = Document()
    doc.add_paragraph("[skill:Python] tools")
    doc.save(src)
    assert cli_main([str(src), "--no-skill-tags", "--dry-run"]) == 0
    assert "1 [skill:…] tag(s)" in capsys.readouterr().out
    assert Document(src).paragraphs[0].text == "[skill:Python] tools"


def test_cli_dry_run_index_only(tmp_path: Path, capsys):
    src = tmp_path / "in.docx"
    doc = Document()
    doc.add_paragraph("[1.1] Spexture Portfolio — [skill:Java]")
    doc.save(src)
    assert cli_main([str(src), "--no-index-tags", "--dry-run"]) == 0
    out = capsys.readouterr().out
    assert "1 index tag(s)" in out
    assert "[skill:…]" not in out
